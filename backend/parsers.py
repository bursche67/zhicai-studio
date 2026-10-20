"""Local-only document parsing. Original attachments are never executed."""
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from io import BytesIO
import hashlib
import os
import shutil
import subprocess
import tempfile
import zipfile
from bs4 import BeautifulSoup

MAX_FILES = 5000
MAX_BYTES = 200 * 1024 * 1024
SUPPORTED = {'.html', '.htm', '.pdf', '.docx', '.xlsx', '.doc', '.png', '.jpg', '.jpeg', '.tif', '.tiff'}

@dataclass
class Block:
    source: str
    locator: str
    text: str

@dataclass
class Parsed:
    blocks: list[Block] = field(default_factory=list)
    tables: list[tuple[str, str, list[list[str]]]] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)

    def extend(self, other):
        self.blocks.extend(other.blocks)
        self.tables.extend(other.tables)
        self.errors.extend(other.errors)

def html_grid(table):
    cells, rows = {}, []
    for r, tr in enumerate(table.find_all('tr')):
        c = 0
        for td in tr.find_all(['td', 'th'], recursive=False):
            while (r, c) in cells:
                c += 1
            text = td.get_text(' ', strip=True)
            rs = min(100, max(1, int(td.get('rowspan', 1))))
            cs = min(100, max(1, int(td.get('colspan', 1))))
            for dr in range(rs):
                for dc in range(cs):
                    cells[r + dr, c + dc] = text
            c += cs
    if cells:
        height = max(r for r, _ in cells) + 1
        width = max(c for _, c in cells) + 1
        rows = [[cells.get((r, c), '') for c in range(width)] for r in range(height)]
    return rows

def ocr(data):
    cmd = os.getenv('TESSERACT_CMD') or shutil.which('tesseract')
    if not cmd:
        raise ValueError('需要OCR：未安装/配置中文Tesseract，内容未提取')
    with tempfile.TemporaryDirectory() as folder:
        image = Path(folder) / 'input.png'
        image.write_bytes(data)
        result = subprocess.run([cmd, str(image), 'stdout', '-l', 'chi_sim+eng'], capture_output=True, timeout=90)
        if result.returncode:
            raise ValueError('OCR失败：请检查chi_sim语言包')
        return result.stdout.decode('utf-8', errors='replace')

def parse_document(name: str, data: bytes) -> Parsed:
    out = Parsed()
    suffix = Path(name).suffix.lower()
    try:
        if suffix in ('.html', '.htm'):
            soup = BeautifulSoup(data, 'html.parser')
            for tag in soup(['script', 'style', 'noscript']):
                tag.decompose()
            # 公告概要含采购单位、完整品目和总金额，不能只取正文容器。
            content = soup.select_one('.vF_detail_main') or soup.select_one('.vF_detail_content_container') or soup.body or soup
            for idx, table in enumerate(content.find_all('table'), 1):
                rows = html_grid(table)
                out.tables.append((name, f'table:{idx}', rows))
            out.blocks.append(Block(name, 'body', content.get_text('\n', strip=True)))
        elif suffix == '.docx':
            from docx import Document
            doc = Document(BytesIO(data))
            for i, p in enumerate(doc.paragraphs, 1):
                if p.text.strip():
                    out.blocks.append(Block(name, f'paragraph:{i}', p.text))
            for i, table in enumerate(doc.tables, 1):
                out.tables.append((name, f'table:{i}', [[c.text for c in row.cells] for row in table.rows]))
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(data), data_only=True)
            for sheet in wb:
                if sheet.max_row * sheet.max_column > 200000:
                    raise ValueError('工作表过大，超过200000单元格限制')
                merged = {}
                for region in sheet.merged_cells.ranges:
                    value = sheet.cell(region.min_row, region.min_col).value
                    for row in sheet.iter_rows(min_row=region.min_row, max_row=region.max_row, min_col=region.min_col, max_col=region.max_col):
                        for c in row:
                            merged[c.coordinate] = value
                rows = [[str(merged.get(c.coordinate, c.value) if merged.get(c.coordinate, c.value) is not None else '') for c in row] for row in sheet]
                out.tables.append((name, f'sheet:{sheet.title}', rows))
            wb.close()
        elif suffix == '.pdf':
            import pymupdf
            with pymupdf.open(stream=data, filetype='pdf') as doc:
                for i, page in enumerate(doc, 1):
                    try:
                        text = page.get_text()
                        if len(text.strip()) < 10:
                            text = ocr(page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png'))
                        out.blocks.append(Block(name, f'page:{i}', text))
                        for t, table in enumerate(page.find_tables().tables, 1):
                            out.tables.append((name, f'page:{i}/table:{t}', [[c or '' for c in row] for row in table.extract()]))
                    except Exception as exc:
                        out.errors.append({'source': name, 'locator': f'page:{i}', 'error': str(exc)})
        elif suffix == '.doc':
            cmd = os.getenv('SOFFICE_PATH') or shutil.which('soffice')
            if not cmd:
                raise ValueError('DOC兼容转换未启用：需要LibreOffice及SOFFICE_PATH')
            with tempfile.TemporaryDirectory() as folder:
                src = Path(folder) / 'input.doc'
                src.write_bytes(data)
                result = subprocess.run([cmd, '-env:UserInstallation=' + (Path(folder) / 'profile').as_uri(), '--headless', '--convert-to', 'docx', '--outdir', folder, str(src)], capture_output=True, timeout=90)
                target = src.with_suffix('.docx')
                if result.returncode or not target.exists():
                    raise ValueError('LibreOffice转换失败')
                out.extend(parse_document(name + '.docx', target.read_bytes()))
        elif suffix in ('.png', '.jpg', '.jpeg', '.tif', '.tiff'):
            out.blocks.append(Block(name, 'ocr', ocr(data)))
        else:
            raise ValueError('暂不支持的附件格式：' + suffix)
        for source, locator, rows in out.tables:
            out.blocks.append(Block(source, locator, '\n'.join(' | '.join(row) for row in rows)))
        out.blocks = [b for b in out.blocks if b.text.strip()]
        if not out.blocks and not out.errors:
            raise ValueError('未提取到可读文本；可能为空文件或公式缓存缺失')
    except Exception as exc:
        out.errors.append({'source': name, 'error': str(exc)})
    return out

def unpack_zip(name: str, data: bytes, depth=0, budget=None):
    """Yield virtual paths and bytes without extracting to the filesystem."""
    if depth > 2:
        raise ValueError('ZIP嵌套超过2层')
    budget = budget if budget is not None else {'files': 0, 'bytes': 0}
    with zipfile.ZipFile(BytesIO(data)) as archive:
        for info in archive.infolist():
            path = PurePosixPath(info.filename.replace('\\', '/'))
            if path.is_absolute() or '..' in path.parts or ':' in str(path):
                raise ValueError('ZIP存在不安全路径')
            if info.is_dir():
                continue
            budget['files'] += 1
            budget['bytes'] += info.file_size
            if budget['files'] > MAX_FILES or budget['bytes'] > MAX_BYTES:
                raise ValueError('ZIP文件数量或展开大小超过限制')
            if info.flag_bits & 1:
                raise ValueError('暂不支持加密ZIP')
            if info.file_size > 1024 * max(info.compress_size, 1):
                raise ValueError('ZIP压缩比超过安全限制')
            content = archive.read(info)
            virtual = name + '/' + str(path)
            if path.suffix.lower() == '.zip':
                yield from unpack_zip(virtual, content, depth + 1, budget)
            else:
                yield virtual, content

def sha(data):
    return hashlib.sha256(data).hexdigest()
