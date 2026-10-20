"""Standard three-sheet export. Application output uses the requested openpyxl stack."""
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from copy import copy
import json
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment
from .db import connection
from .models import Extraction

HEADERS = {
    '总表': ['序号', 'URL', 'UUID', '标段/合同包', '项目名称', '采购单位', '中标供应商', '中标金额'],
    '标的表': ['序号', 'UUID', '标段/合同包', '产品服务名称', '品目', '品牌', '规格型号', '数量', '单价', '总价', '类别'],
    '投标人表': ['序号', 'UUID', '标段/合同包', '名称', '类别', '通过审查', '报价（万元）', '得分', '未中标原因'],
}
TEMPLATE = Path(__file__).parent / 'resources' / 'standard_output_template.xlsx'

def wan(value, rounded=False):
    if value is None:
        return None
    v = Decimal(value) / Decimal(10000)
    if rounded:
        v = v.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    return float(v)

def safe(value):
    # Force untrusted source strings to remain text in spreadsheet readers.
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value

def latest_documents(mode, sample=False, job_id=None):
    with connection() as db:
        sql = '''SELECT d.* FROM documents d JOIN jobs j ON j.id=d.job_id WHERE j.mode=? AND d.is_sample=?
        AND d.status IN ('completed','partial','duplicate') AND d.extraction IS NOT NULL'''
        args = [mode, int(sample)]
        if job_id:
            sql += ' AND d.job_id=?'
            args.append(job_id)
        docs = db.execute(sql + ' ORDER BY j.created_at DESC,d.id DESC', args).fetchall()
    seen, selected = set(), []
    for doc in docs:
        key = doc['source_uuid'] or 'unidentified:' + doc['id']
        if key not in seen:
            seen.add(key)
            selected.append(doc)
    return selected


def standard_rows(mode, sample=False, job_id=None):
    tables = {name: [] for name in HEADERS}
    warnings = []
    docs = latest_documents(mode, sample, job_id)
    seen = set()
    for doc in docs:
        # Source UUID identifies the notice, independent of ZIP/path naming.
        # Select the latest successful version, not repeated historical runs.
        key = doc['source_uuid']
        if not key:
            warnings.append(f'{doc["name"]}: 缺少公告UUID')
            continue
        if key in seen:
            continue
        seen.add(key)
        out = Extraction.model_validate_json(doc['extraction'])
        uuid = doc['source_uuid']
        if not uuid:
            warnings.append(f'{doc["name"]}: 缺少公告UUID')
            continue
        for p in out.packages:
            package_name = p.name or p.code
            # Only explicit unsegmented notices may say 未分包. Unknown remains blank.
            winners = [b for b in p.bids if b.result == 'won']
            winner_units = {b.consortium or b.supplier for b in winners}
            if winners:
                seen_groups = set()
                for b in winners:
                    if b.consortium and b.consortium in seen_groups:
                        continue
                    if b.consortium:
                        seen_groups.add(b.consortium)
                    amount = b.consortium_award_amount if b.consortium else b.award_amount
                    if amount is None and len(winner_units) == 1:
                        amount = p.award_amount
                    if amount is None and len(out.packages) == 1 and len(winner_units) == 1:
                        amount = out.award_amount
                    tables['总表'].append([None, doc['source_url'], uuid, package_name, out.project_name, out.buyer, b.consortium or b.supplier, wan(amount, True)])
            else:
                tables['总表'].append([None, doc['source_url'], uuid, package_name, out.project_name, out.buyer, None, wan(p.award_amount, True)])
            for i in p.items:
                total = i.total_price
                if total is None and i.unit_price is not None and i.quantity is not None:
                    total = i.unit_price * i.quantity
                    warnings.append(f'{uuid}/{package_name}/{i.name}：导出总价由单价×数量计算，非原文披露；数据库原始字段仍为null')
                quantity = None if i.quantity is None else int(i.quantity) if i.quantity == i.quantity.to_integral_value() else float(i.quantity)
                tables['标的表'].append([None, uuid, package_name, i.name, i.category, i.brand, i.specification, quantity, wan(i.unit_price), wan(total), None])
            seen_groups = set()
            for b in p.bids:
                if b.consortium and b.consortium in seen_groups:
                    continue
                if b.consortium:
                    seen_groups.add(b.consortium)
                category = '中标人' if b.result == 'won' else '其他' if b.result in ('invalid', 'cancelled') else '参与人'
                score = b.consortium_score if b.consortium else b.score
                review = b.consortium_review_status if b.consortium else b.review_status
                price = b.consortium_bid_price if b.consortium else b.bid_price
                tables['投标人表'].append([None, uuid, package_name, b.consortium or b.supplier, category, review, wan(price, True), float(score.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)) if score is not None else None, b.loss_reason])
        if not out.packages:
            tables['总表'].append([None, doc['source_url'], uuid, None, out.project_name, out.buyer, None, wan(out.award_amount, True)])
        warnings.extend(out.warnings)
    for rows in tables.values():
        for index, row in enumerate(rows, 1):
            row[0] = index
    return tables, warnings

def workbook_bytes(mode='model', sample=False, job_id=None):
    tables, _ = standard_rows(mode, sample, job_id)
    wb = load_workbook(TEMPLATE)
    for name, headers in HEADERS.items():
        sheet = wb[name]
        if [c.value for c in sheet[1]] != headers:
            raise ValueError('标准输出模板表头发生变化，需重新核验导出映射')
        instructions = [c.value for c in sheet[2]]
        data_styles = [copy(c._style) for c in sheet[4]]
        # Exclude the explanatory/example rows from exported results.
        # Retain their instructions as header comments and their native styling.
        sheet.delete_rows(2, sheet.max_row - 1)
        for cell, note in zip(sheet[1], instructions):
            if note:
                cell.comment = Comment(str(note), '字段填写说明')
        for row in tables[name]:
            sheet.append([safe(v) for v in row])
            for cell, style in zip(sheet[sheet.max_row], data_styles):
                cell._style = copy(style)
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '0.########'
        if name == '总表':
            for cell in list(sheet.columns)[7][1:]:
                cell.number_format = '0.00'
        if name == '投标人表':
            for col in (7, 8):
                for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                    row[0].number_format = '0.00'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
