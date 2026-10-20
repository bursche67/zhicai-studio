from io import BytesIO
from copy import copy
from openpyxl import load_workbook
from backend.export import workbook_bytes, standard_rows, TEMPLATE, HEADERS
from backend.jobs import create_job, process_job
from test_core import HTML


def test_template_headers_styles_units_and_no_examples(database):
    job = create_job([('t20260101_123.html', HTML)], 'rule', True)
    process_job(job)
    result = load_workbook(BytesIO(workbook_bytes('rule', True)))
    template = load_workbook(TEMPLATE)
    assert result.sheetnames == template.sheetnames
    for name in HEADERS:
        assert [c.value for c in result[name][1]] == [c.value for c in template[name][1]]
        for a, b in zip(result[name][1], template[name][1]):
            assert copy(a.fill) == copy(b.fill)
            assert copy(a.font) == copy(b.font)
            assert a.comment.text == str(template[name].cell(2, a.column).value)
        assert not any('填写示例' in str(c.value) or '示例数据' in str(c.value) for row in result[name] for c in row)
    assert result['标的表'].cell(2, 8).data_type == 'n'
    assert result['标的表'].cell(2, 8).value == 2
    assert result['标的表'].cell(2, 9).value == 0.5  # 5000 CNY -> 0.5万元
    assert result['总表'].cell(2, 8).number_format == '0.00'


def test_latest_successful_notice_version_only(database):
    old = create_job([('t20260101_123.html', HTML)], 'rule', True)
    process_job(old)
    new = create_job([('t20260101_123.html', HTML.replace('2台'.encode(), '3台'.encode()))], 'rule', True)
    process_job(new)
    tables, _ = standard_rows('rule', True)
    assert len(tables['标的表']) == 1
    assert tables['标的表'][0][7] == 3
    old_tables, _ = standard_rows('rule', True, old)
    assert old_tables['标的表'][0][7] == 2
