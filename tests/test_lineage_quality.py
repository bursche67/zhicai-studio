from pathlib import Path
import zipfile
from io import BytesIO
from openpyxl import load_workbook
from backend.models import Extraction, Package, Bid, Item
from backend.parsers import parse_document
from backend.extraction import source_guard
from backend.jobs import create_job,process_job
from backend.db import connection, persist
from backend.export import workbook_bytes
from backend.analysis import cooperation, buyer_bidders, common_projects, supplier_competitors, common_buyers
import pytest


def store(out, doc, uid):
    jid = create_job([(uid+'.html', b'<body>test fixture</body>')], 'model')
    with connection() as db:
        db.execute('INSERT INTO documents(id,job_id,name,sha256,status,source_uuid) VALUES(?,?,?,?,?,?)', (doc,jid,uid+'.html',doc,'completed',uid))
        return persist(db,doc,out,'model')


def test_same_uuid_new_version_replaces_relations_and_older_replay_cannot_restore(database):
    old = Extraction(buyer='采购人',project_number='P',packages=[Package(name='甲包',bids=[Bid(supplier='公司甲',result='won',award_amount=100)])])
    new = old.model_copy(deep=True)
    new.packages[0].bids[0].loss_reason = '原文审查原因'
    pid = store(old,'old','same-uuid')
    store(new,'new','same-uuid')
    with connection() as db:
        persist(db,'old',old,'model')
        rows=db.execute('SELECT * FROM bids WHERE project_id=?',(pid,)).fetchall()
        assert len(rows)==1 and rows[0]['loss_reason']=='原文审查原因'
    assert cooperation('采购人')['winners'][0]['known_amount_cents']==10000


def test_same_package_duplicate_keeps_non_null_field(database):
    out=Extraction(buyer='采购人',project_number='P',packages=[Package(code='1',bids=[
        Bid(supplier='公司甲',result='invalid',loss_reason='资格性未通过'),
        Bid(supplier='公司甲',result='invalid')])])
    pid=store(out,'doc','u')
    with connection() as db:
        assert db.execute('SELECT loss_reason FROM bids WHERE project_id=?',(pid,)).fetchone()[0]=='资格性未通过'


@pytest.mark.parametrize('reverse',[False,True])
def test_conflicting_win_and_loss_validates_only_after_clearing_award(database,reverse):
    bids=[Bid(supplier='公司甲',result='lost',loss_reason='未通过'),Bid(supplier='公司甲',result='won',award_amount=100)]
    if reverse:bids.reverse()
    pid=store(Extraction(buyer='采购人',project_number='P',packages=[Package(code='1',bids=bids)]),'doc','u')
    with connection() as db:
        row=db.execute('SELECT result,award_amount,loss_reason FROM bids WHERE project_id=?',(pid,)).fetchone()
        assert row['result']=='unknown' and row['award_amount'] is None and row['loss_reason']=='未通过'


@pytest.mark.parametrize('reverse',[False,True])
def test_explicit_invalid_refines_generic_non_award(database,reverse):
    bids=[Bid(supplier='公司甲',result='invalid',loss_reason='符合性审查未通过'),Bid(supplier='公司甲',result='lost')]
    if reverse:bids.reverse()
    pid=store(Extraction(buyer='采购人',project_number='P',packages=[Package(code='1',bids=bids)]),'doc','u')
    with connection() as db:
        row=db.execute('SELECT result,loss_reason FROM bids WHERE project_id=?',(pid,)).fetchone()
        assert row['result']=='invalid' and row['loss_reason']=='符合性审查未通过'


def test_award_is_not_bid_price_and_wrong_package_scores_not_filled():
    p=parse_document('u.html','<body>第3包中标（成交）金额：100元。第1包中标人评审得分：96.00</body>'.encode())
    x=Extraction(packages=[Package(code='3',bids=[Bid(supplier='公司甲',result='won',award_amount=100,bid_price=100)])])
    r=source_guard(x,p)
    assert r.packages[0].bids[0].bid_price is None
    assert r.packages[0].bids[0].score is None
    p=parse_document('u.html','<body>公司甲最终报价：100元</body>'.encode())
    x.packages[0].bids[0].bid_price=100
    assert source_guard(x,p).packages[0].bids[0].bid_price==100
    p=parse_document('u.html','<body>中标金额100元，附件：报价明细附件.pdf</body>'.encode())
    x.packages[0].bids[0].bid_price=100
    assert source_guard(x,p).packages[0].bids[0].bid_price is None


def test_consortium_members_not_a_pseudo_company_and_group_money_once(database):
    text='<table><tr><th>供应商</th><th>中标金额</th></tr><tr><td>牵头供应商：公司甲有限责任公司 投标联合体：公司甲有限责任公司、公司乙有限责任公司</td><td>100元</td></tr></table>'
    parsed=parse_document('u.html',text.encode())
    x=Extraction(buyer='采购人',project_number='P',project_amount=120,award_amount=100,packages=[Package(code='1',award_amount=100,bids=[Bid(supplier='公司甲有限责任公司、公司乙有限责任公司（联合体）',result='won',award_amount=100,review_status='通过',score=90)])])
    result=source_guard(x,parsed)
    assert {b.supplier for b in result.packages[0].bids}=={'公司甲有限责任公司','公司乙有限责任公司'}
    assert all(b.award_amount is None and b.score is None and b.consortium_score==90 for b in result.packages[0].bids)
    store(result,'doc','u')
    assert len(cooperation('采购人')['winners'])==2
    assert all(w['known_amount_cents'] is None for w in cooperation('采购人')['winners'])
    names=['公司甲有限责任公司','公司乙有限责任公司']
    assert len(buyer_bidders('采购人')['pairs'])==1
    assert supplier_competitors(names[0])['all'][0]['supplier']==names[1]
    assert common_buyers(names)['buyers'][0]['project_count']==1
    assert common_projects(names)['project_count']==1
    book=load_workbook(BytesIO(workbook_bytes()))
    assert book['总表'].max_row==2 and book['总表'].cell(2,8).value==0.01
    assert book['投标人表'].max_row==2 and book['投标人表'].cell(2,8).value==90


def test_product_supplier_role_does_not_implicitly_create_bid(database):
    result=Extraction(buyer='采购人',project_number='P',packages=[Package(code='1',
        bids=[Bid(supplier='实际中标公司',result='won'),Bid(supplier='实际参与公司',result='lost')],
        items=[Item(name='设备',brand='仅品牌文字',product_supplier='明确供货公司')])])
    pid=store(result,'roles','u')
    with connection() as db:
        bids={r[0] for r in db.execute('SELECT o.name FROM bids b JOIN organizations o ON o.id=b.supplier_id WHERE b.project_id=?',(pid,))}
        assert bids=={'实际中标公司','实际参与公司'}
        assert db.execute('SELECT o.name FROM items i JOIN organizations o ON o.id=i.product_supplier_id WHERE i.project_id=?',(pid,)).fetchone()[0]=='明确供货公司'


def test_duplicate_reuses_latest_successful_extraction(database,monkeypatch):
    from backend.parsers import sha
    old=Extraction(buyer='采购人',project_number='P',project_name='旧结果')
    new=old.model_copy(update={'project_name':'最新结果'})
    store(old,'old-doc','u');store(new,'new-doc','u')
    digest=sha(b'u.html\0<body>test fixture</body>')
    with connection() as db:db.execute('UPDATE documents SET sha256=?',(digest,))
    monkeypatch.setattr('backend.jobs.model_unavailable_reason',lambda:None)
    def forbidden(*args):raise AssertionError('Duplicate must make zero model calls')
    monkeypatch.setattr('backend.jobs.model_extract',forbidden)
    jid=create_job([('u.html',b'<body>test fixture</body>')],'model')
    process_job(jid)
    with connection() as db:
        row=db.execute('SELECT status,extraction FROM documents WHERE job_id=?',(jid,)).fetchone()
        assert row['status']=='duplicate'
        assert Extraction.model_validate_json(row['extraction']).project_name=='最新结果'


def test_actual_target_html_retains_score_package_mismatch():
    dataset=next(Path('.').glob('*_html.zip'),None)
    if dataset is None:
        import pytest
        pytest.skip('真实数据回归需要可选HTML回归数据集；测试源码不捆绑该数据集')
    with zipfile.ZipFile(dataset) as z:
        name=next(n for n in z.namelist() if Path(n).stem=='t20260202_26139731')
        parsed=parse_document(name,z.read(name))
    body=next(b.text for b in parsed.blocks if b.locator=='body')
    assert not parsed.errors
    assert '第3包 总控光配线系统：北京万迪科信息技术有限公司' in body
    assert '第4包 LCD显示大屏（第一部分）：北京智为视媒科技有限公司' in body
    assert '第1包中标人评审得分：96.00' in body
    assert '第2包中标人评审得分：97.93' in body
    result=source_guard(Extraction(packages=[Package(code='第3包',bids=[Bid(supplier='北京万迪科信息技术有限公司',result='won')],items=[Item(name='光跳纤架',brand='康普'),Item(name='光熔纤柜',brand='康普')])]),parsed)
    assert all(i.product_supplier=='北京万迪科信息技术有限公司' for i in result.packages[0].items)
    assert all(i.brand=='康普' for i in result.packages[0].items)
    assert len(result.packages[0].bids)==1
