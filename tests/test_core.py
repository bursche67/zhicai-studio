from decimal import Decimal
from io import BytesIO
import json
import zipfile
import pytest
import httpx
from openpyxl import load_workbook, Workbook
from docx import Document
from backend.models import money, Extraction, Package, Item, Bid
from backend.parsers import parse_document, unpack_zip
from backend.extraction import rule_extract, model_extract, merge_extractions, source_guard
from backend.jobs import prepare_groups, create_job, process_job
from backend.db import connection, persist
from backend.export import workbook_bytes, standard_rows
from backend.analysis import cooperation, buyer_bidders, supplier_competitors, common_buyers, common_projects
from backend.graph import graph_adapter

HTML = '''<html><link rel="canonical" href="https://www.ccgp.gov.cn/test/t20260101_123.html"><body><p>项目名称：自建测试项目</p><p>项目编号：SAMPLE-1</p><p>采购单位：自建测试大学</p><p>项目金额：2万元</p><p>中标金额：1万元</p><p>中标供应商：样例甲</p><p>投标参与方：样例乙</p><table><tr><th>产品服务名称</th><th>品目</th><th>品牌</th><th>规格型号</th><th>单价</th><th>数量</th><th>总价</th></tr><tr><td>样例服务器</td><td>设备</td><td>样例品牌</td><td>S-1</td><td>5000元</td><td>2台</td><td>10000元</td></tr></table></body></html>'''.encode()

@pytest.mark.parametrize('raw,expected', [('1.25万元', '12500.00'), ('￥1,234.565（人民币）', '1234.57'), (None, None), ('0','0.00')])
def test_money(raw, expected):
    v = money(raw)
    assert (str(v) if v is not None else None) == expected

@pytest.mark.parametrize('raw', ['-1','NaN','Infinity','12美元','1至3万元',True])
def test_money_rejects(raw):
    with pytest.raises(ValueError):
        money(raw)

def test_html_rule_and_merged_cells():
    parsed = parse_document('sample.html', HTML)
    out = rule_extract(parsed)
    assert out.buyer == '自建测试大学'
    item = out.packages[0].items[0]
    assert item.quantity == 2 and item.unit == '台' and item.total_price == 10000
    assert item.evidence[0].locator.endswith('row:2')
    parsed = parse_document('merged.html', b'<table><tr><td rowspan="2">A</td><td>B</td></tr><tr><td>C</td></tr></table>')
    assert parsed.tables[0][2] == [['A','B'],['A','C']]
    parsed = parse_document('category.html','<table><tr><th>品目号</th><th>采购标的</th></tr><tr><td>1-1</td><td>设备</td></tr></table>'.encode())
    assert rule_extract(parsed).packages[0].items[0].category is None
    html='<div class="vF_detail_main"><table><tr><td>采购单位</td><td>保留概要采购人</td></tr></table><div class="vF_detail_content_container">正文</div></div>'
    assert '保留概要采购人' in parse_document('summary.html',html.encode()).blocks[0].text
    out=source_guard(Extraction(packages=[Package(name='未分包')]),parsed)
    assert out.packages[0].name is None and out.warnings

def test_docx_xlsx_pdf():
    doc = Document();doc.add_paragraph('项目名称：测试')
    table = doc.add_table(rows=1, cols=2);table.cell(0,0).text='产品服务名称';table.cell(0,1).text='数量'
    stream=BytesIO();doc.save(stream)
    assert not parse_document('a.docx', stream.getvalue()).errors
    book=Workbook();sheet=book.active;sheet.append(['A','B']);sheet.append(['same',1]);sheet.append([None,2]);sheet.merge_cells('A2:A3')
    stream=BytesIO();book.save(stream)
    parsed = parse_document('a.xlsx',stream.getvalue())
    assert parsed.tables[0][2][2][0] == 'same'
    import pymupdf
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((72,72),'sample procurement text')
    assert 'sample procurement' in parse_document('a.pdf',pdf.tobytes()).blocks[0].text

def test_unsupported_and_corrupt(monkeypatch):
    monkeypatch.setattr('backend.parsers.shutil.which',lambda _:None)
    monkeypatch.delenv('SOFFICE_PATH',raising=False);monkeypatch.delenv('TESSERACT_CMD',raising=False)
    assert 'DOC' in parse_document('a.doc',b'bad').errors[0]['error']
    assert 'OCR' in parse_document('a.png',b'bad').errors[0]['error']
    assert parse_document('bad.pdf',b'bad').errors

def archive(entries):
    b=BytesIO()
    with zipfile.ZipFile(b,'w') as z:
        for n,data in entries:
            z.writestr(n,data)
    return b.getvalue()

def test_zip_security():
    with pytest.raises(ValueError, match='不安全'):
        list(unpack_zip('x.zip',archive([('../outside',b'bad')])))
    with pytest.raises(ValueError, match='不安全'):
        list(unpack_zip('x.zip',archive([('C:/bad',b'bad')])))
    nested=archive([('file.txt',b'ok')])
    for _ in range(3):
        nested=archive([('inner.zip',nested)])
    with pytest.raises(ValueError, match='嵌套'):
        list(unpack_zip('x.zip',nested))

def test_pairing():
    files,groups,issues=prepare_groups([('all.zip',archive([('a.html',HTML),('b.html',HTML)])),('a.zip',archive([('x.pdf',b'bad')])),('unmatched.doc',b'bad')])
    assert len(groups)==2 and len(groups[0])==2
    assert issues[0]['name']=='unmatched.doc'
    with pytest.raises(ValueError):
        prepare_groups([('a.html',HTML)],{'groups':[{'html':'missing.html'}]})

def insert(db, did, out):
    db.execute('INSERT OR IGNORE INTO jobs(id,mode,status,created_at,manifest) VALUES(?,?,?,?,?)',('fixtures','model','completed','2026-01-01','{}'))
    db.execute('INSERT INTO documents(id,job_id,name,sha256,status,is_sample,source_uuid,extraction) VALUES(?,?,?,?,?,?,?,?)',(did,'fixtures',did+'.html',did,'completed',1,did,out.model_dump_json()))
    return persist(db,did,out,'model',True)

@pytest.fixture
def relation_data(database):
    def project(number,buyer,packages,amount='100',award='80'):
        return Extraction(project_name=number,project_number=number,buyer=buyer,project_amount=amount,award_amount=award,packages=packages)
    def p(code,bids,items=None):
        return Package(code=code,bids=[Bid(supplier=n,result=r,award_amount=a) for n,r,a in bids],items=items or [])
    with connection() as db:
        a=insert(db,'source-1',project('P1','采购方一',[p('1',[('甲','won','60'),('乙','lost',None),('丙','lost',None)],[Item(name='商品',product_supplier='厂商',total_price=20)]),p('2',[('甲','won','20'),('乙','lost',None)])]))
        insert(db,'source-2',project('P2','采购方一',[p('1',[('乙','won','80'),('甲','lost',None)])],amount=None))
        insert(db,'source-3',project('P3','采购方二',[p('1',[('甲','won',None),('乙','won',None)])],amount='200',award='120'))
        insert(db,'source-4',project('P4','采购方三',[p('1',[('甲','won','10')]),p('2',[('乙','won','20')])],amount='300',award='30'))
    return a

def test_five_scenarios(relation_data):
    co=cooperation('采购方一','model',True)
    assert co['winners'][0]['cooperation_count']==1
    assert {r['supplier']:r['known_amount_cents'] for r in co['winners']}=={'甲':8000,'乙':8000}
    assert co['product_suppliers'][0]['known_amount_cents']==2000
    bidders=buyer_bidders('采购方一','model',True)
    assert [r['frequency'] for r in bidders['top5']]==[2,2,1]
    assert max(r['frequency'] for r in bidders['pairs'])==2
    competitors=supplier_competitors('甲','model',True)
    assert competitors['top5'][0]['supplier']=='乙' and competitors['top5'][0]['frequency']==4
    assert competitors['top5'][0]['missing_amount_count']==1
    buyers=common_buyers(['甲','乙'],'model',True)['buyers']
    assert {r['buyer'] for r in buyers}=={'采购方一','采购方二','采购方三'}
    assert next(r for r in buyers if r['buyer']=='采购方一')['project_count']==2
    assert next(r for r in buyers if r['buyer']=='采购方二')['known_amount_cents']==12000
    common=common_projects(['甲','乙'],'model',True)
    assert common['project_count']==4 and common['known_amount_cents']==60000 and common['missing_amount_count']==1
    assert common_projects(['甲','丙','乙'],'model',True)['project_count']==1
    assert common_projects(['甲','不存在'],'model',True)['project_count']==0
    assert cooperation('采购方一','rule',True)['winners']==[]

def test_rerun_and_conflict(relation_data):
    with connection() as db:
        out=Extraction(project_name='P1',project_number='P1',buyer='采购方一',packages=[Package(code='1',bids=[Bid(supplier='甲',result='won',award_amount=10)])])
        persist(db,'source-1',out,'model',True)
    co=cooperation('采购方一','model',True)
    assert next(r for r in co['winners'] if r['supplier']=='甲')['known_amount_cents']==1000

def test_graph(relation_data):
    graph=graph_adapter.get_graph('model',True,organization='甲',limit=10)
    assert graph['nodes'] and len(graph['nodes'])<=10
    ids={n['id'] for n in graph['nodes']}
    assert all(e['source'] in ids and e['target'] in ids for e in graph['links'])

def test_model_missing_config():
    import os
    old=os.environ.pop('MODEL_NAME',None)
    try:
        with pytest.raises(ValueError,match='配置缺失'):
            model_extract(parse_document('a.html',HTML),[])
    finally:
        if old:os.environ['MODEL_NAME']=old

def test_model_mock_valid_invalid_timeout(monkeypatch):
    monkeypatch.setenv('MODEL_BASE_URL','http://localhost:8001/v1');monkeypatch.setenv('MODEL_NAME','Qwen-test');monkeypatch.setenv('MODEL_RETRIES','0')
    parsed=parse_document('a.html',HTML)
    body=Extraction(project_name='自建测试项目').model_dump_json()
    monkeypatch.setenv('MODEL_THINKING','disabled')
    monkeypatch.setenv('MODEL_JSON_MODE','true')
    def response(self,*a,**kw):
        assert kw['json']['thinking']=={'type':'disabled'}
        assert kw['json']['response_format']=={'type':'json_object'}
        assert all(b['locator']=='body' for b in json.loads(kw['json']['messages'][1]['content']))
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':body}}]})
    monkeypatch.setattr(httpx.Client,'post',response)
    trace=[];assert model_extract(parsed,trace).project_name=='自建测试项目' and trace[0]['status']=='validated'
    body='not json'
    with pytest.raises(ValueError,match='合法JSON'):model_extract(parsed,[])
    def timeout(*a,**kw):raise httpx.ReadTimeout('test')
    monkeypatch.setattr(httpx.Client,'post',timeout)
    with pytest.raises(ValueError,match='超时'):model_extract(parsed,[])

def test_import_export_and_duplicate(database):
    job=create_job([('t20260101_123.html',HTML)],'rule',True);process_job(job)
    with connection() as db:
        assert db.execute('SELECT status FROM jobs WHERE id=?',(job,)).fetchone()[0]=='completed'
        assert db.execute('SELECT count(*) FROM items').fetchone()[0]==1
    duplicate=create_job([('t20260101_123.html',HTML)],'rule',True);process_job(duplicate)
    with connection() as db:
        assert db.execute('SELECT status FROM documents WHERE job_id=?',(duplicate,)).fetchone()[0]=='duplicate'
        assert db.execute('SELECT count(*) FROM items').fetchone()[0]==1
    book=load_workbook(BytesIO(workbook_bytes('rule',True)))
    assert book.sheetnames==['总表','标的表','投标人表']
    assert book['总表'].cell(2,3).value=='t20260101_123'
    assert book['总表'].cell(2,8).value==1
    assert book['标的表'].cell(2,9).value==.5
    assert book['投标人表'].max_row==3
    # Retry replaces same rows rather than appending.
    process_job(job)
    with connection() as db:
        assert db.execute('SELECT count(*) FROM items').fetchone()[0]==1

def test_failed_model_does_not_write_items(database):
    job=create_job([('sample.html',HTML)],'model',True);process_job(job)
    with connection() as db:
        assert db.execute('SELECT status FROM jobs WHERE id=?',(job,)).fetchone()[0]=='failed'
        assert db.execute('SELECT count(*) FROM items').fetchone()[0]==0
